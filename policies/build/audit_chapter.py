# -*- coding: utf-8 -*-
"""Chapter-wide SOP/Record audit -- the same checks run against COP on
2026-10-10 (which caught the COP.2 triage/reassessment bug), generalised
to any chapter folder under sop_record_masters/.

Checks:
  1. Unsourced numeric "e.g." examples (digit inside an e.g.,... parenthetical)
     with no "Technical source(s)" anywhere in that heading block.
  2. Named-scale/category examples presented as a slash-separated list
     inside an "e.g." parenthetical (the "Emergency / Priority / Non-urgent"
     pattern) -- flagged for manual look regardless of sourcing, since even
     a sourced one is worth a human glance.
  3. Every bare numeric time/duration value anywhere in the file, with
     context, for manual review (most will be fine -- statutory citations,
     or illustrative Record worked-example data -- but each needs eyes).
  4. Known artifact-bug strings (single-brace {HOSPITAL_NAME}, the
     "Guidebook interpretation supplied" build-note leak, dead "refer to
     the Glossary" pointers, unbalanced guillemets).

This script does not judge sourcing of the time values it lists under (3)
-- that judgment call is made by reading each hit's context, same as was
done for COP by hand. It narrows what needs a human look; it doesn't
replace the look.

Usage:
  python policies/build/audit_chapter.py AAC
  python policies/build/audit_chapter.py ROM
  python policies/build/audit_chapter.py HRM
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).parent / "sop_record_masters"

TIME_RE = re.compile(r'\b\d+[\d\-–\s]*(?:minute|min|hour|hr|day|week|month|year)s?\b', re.I)
EG_DIGIT_RE = re.compile(r'e\.g\.,[^)]*\d')
EG_SLASH_RE = re.compile(r'\(e\.g\.,[^)]*/[^)]*\)')


def full_text(path: Path) -> str:
    d = Document(str(path))
    chunks = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                chunks.append(c.text)
    return "\n".join(chunks)


def sop_headings_with_blocks(path: Path) -> list[tuple[str, str]]:
    d = Document(str(path))
    paras = [p.text for p in d.paragraphs]
    idxs = [i for i, t in enumerate(paras) if re.match(r"4\.\d+ ", t)]
    idxs.append(len(paras))
    out = []
    for i in range(len(idxs) - 1):
        start, end = idxs[i], idxs[i + 1]
        out.append((paras[start], " ".join(paras[start:end])))
    return out


def audit(chapter: str) -> None:
    base = ROOT / chapter
    sop_files = sorted(base.glob("SOP_*.docx"))
    record_files = sorted((base / "Records").glob("*.docx")) if (base / "Records").exists() else []
    print(f"=== {chapter}: {len(sop_files)} SOP(s), {len(record_files)} Record(s) ===\n")

    print("--- Check 1+2: unsourced/slash e.g. examples in SOPs ---")
    any_hit = False
    for f in sop_files:
        for heading, block in sop_headings_with_blocks(f):
            if EG_DIGIT_RE.search(block):
                has_src = "Technical source" in block
                print(f"  [{f.name}] {heading} -> e.g.+digit, Technical source present: {has_src}")
                any_hit = True
            if EG_SLASH_RE.search(block):
                print(f"  [{f.name}] {heading} -> slash-list e.g. example (check if it names a real scale)")
                any_hit = True
    if not any_hit:
        print("  none found")

    print("\n--- Check 3: all bare numeric time/duration values (context, for manual review) ---")
    for f in sop_files:
        t = full_text(f)
        for m in TIME_RE.finditer(t):
            print(f"  [SOP:{f.name}] {m.group(0)!r} :: ...{t[max(0,m.start()-70):m.end()+30]!r}...")
    for f in record_files:
        t = full_text(f)
        for m in TIME_RE.finditer(t):
            print(f"  [REC:{f.name}] {m.group(0)!r} :: ...{t[max(0,m.start()-70):m.end()+30]!r}...")

    print("\n--- Check 4: known artifact bugs ---")
    clean = True
    for f in sop_files + record_files:
        t = full_text(f)
        issues = []
        if re.search(r'(?<!{){HOSPITAL_NAME}(?!})', t):
            issues.append("single-brace {HOSPITAL_NAME}")
        if "Guidebook interpretation supplied" in t:
            issues.append("internal build-note artifact")
        if re.search(r'refer to the glossary', t, re.I):
            issues.append("dead glossary pointer")
        if t.count("«") != t.count("»"):
            issues.append(f"unbalanced guillemets ({t.count(chr(171))} open / {t.count(chr(187))} close)")
        if issues:
            print(f"  [{f.name}] -> {issues}")
            clean = False
    if clean:
        print("  none found")
    print()


if __name__ == "__main__":
    for ch in sys.argv[1:]:
        audit(ch)
