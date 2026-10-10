# -*- coding: utf-8 -*-
"""Structural self-check for sop_record_common.py output.

Confirms a freshly built SOP/Record .docx matches the template shape already
shipped and approved for COP (SOP_COP.13.* and its Records), rather than
trusting the builder by construction. Checks structure only -- it has no
opinion on the content, which is a separate, per-standard verification.

Usage:
  python policies/build/verify_sop_record_template.py <new_sop.docx> --kind sop
  python policies/build/verify_sop_record_template.py <new_record.docx> --kind record
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document

EXPECTED_SOP_HEADINGS = [
    "1. Purpose",
    "2. Scope",
    "3. Responsibility",
    "4. Procedure",
    "5. Records",
    "6. Related Documents",
    "8. Review",
]


def shd_fill(cell) -> str | None:
    m = re.search(r'<w:shd[^/]*w:fill="([0-9A-Fa-f]{6})"', cell._tc.xml)
    return m.group(1) if m else None


def verify_sop(path: Path) -> list[str]:
    errors = []
    doc = Document(str(path))
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    for h in EXPECTED_SOP_HEADINGS:
        if not any(h in x for x in headings):
            errors.append(f"missing Heading 1 section containing '{h}'")
    if "Where This SOP's Requirements Come From" not in "\n".join(headings):
        errors.append("missing section 7 (Where This SOP's Requirements Come From)")
    if len(doc.tables) < 2:
        errors.append(f"expected >=2 tables (intro cards + control box), found {len(doc.tables)}")
    for t in doc.tables:
        if t.style is None or t.style.name != "Table Grid":
            errors.append(f"table style is {t.style.name if t.style else None}, expected 'Table Grid'")
    para0 = doc.paragraphs[0].text if doc.paragraphs else ""
    if para0 != "Standard Operating Procedure":
        errors.append(f"first paragraph is {para0!r}, expected 'Standard Operating Procedure'")
    return errors


def verify_record(path: Path) -> list[str]:
    errors = []
    doc = Document(str(path))
    if len(doc.tables) < 4:
        errors.append(f"expected >=4 tables (3 intro cards + data grid), found {len(doc.tables)}")
        return errors
    data_table = doc.tables[-1]
    if data_table.style is None or data_table.style.name != "Table Grid":
        errors.append("data table style is not 'Table Grid'")
    if len(data_table.rows) < 3:
        errors.append("data table has fewer than 3 rows (header + spacer + example)")
        return errors
    spacer_fill = shd_fill(data_table.rows[1].cells[0])
    example_fill = shd_fill(data_table.rows[2].cells[0])
    if spacer_fill != "FFF2CC":
        errors.append(f"spacer row shading is {spacer_fill!r}, expected 'FFF2CC'")
    if example_fill != "FFF2CC":
        errors.append(f"worked-example row shading is {example_fill!r}, expected 'FFF2CC'")
    example_row_text = [c.text for c in data_table.rows[2].cells]
    if all(v == "" for v in example_row_text):
        errors.append("worked-example row (row index 2) is blank -- no actual example data")
    return errors


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--kind", choices=["sop", "record"], required=True)
    args = ap.parse_args()
    path = Path(args.path)
    if not path.exists():
        print(f"FAIL: no such file: {path}")
        return 1
    errors = verify_sop(path) if args.kind == "sop" else verify_record(path)
    if errors:
        print(f"FAIL: {path.name}")
        for e in errors:
            print(f"  - {e}")
        return 1
    print(f"OK: {path.name} matches the approved COP template shape.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
